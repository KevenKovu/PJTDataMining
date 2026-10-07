# -*- coding: utf-8 -*-
"""
AP5 - Análise descritiva (Table One) | PNS 2019 | Colesterol alto | 25 a 59 anos
==============================================================================
Lê o CSV JÁ FILTRADO (25 a 59 anos, questionário do adulto selecionado concluído,
Q060 = "Sim") e gera as duas tabelas pedidas em aula:

  (1) Variáveis numéricas:
      Variável | Tipo (I/R) | Nº de Ausentes | Mínimo | Máximo | Média | Mediana | Desvio Padrão

  (2) Variáveis categóricas:
      Variável | Tipo (B/O/N) | Nº de Ausentes | Valores | Frequência por valor | Moda

Como usar no VS Code
--------------------
1) Coloque este arquivo .py e o CSV (pns_colesterol_25_59_v3.csv) na MESMA pasta.
   Se o CSV estiver em outro lugar, ajuste NOME_CSV abaixo ou passe o caminho
   como argumento:  python ap5_tabela_descritiva_colesterol.py "C:\\caminho\\arquivo.csv"
2) No terminal do VS Code:  pip install pandas openpyxl
3) Clique em "Run Python File" (▶) ou rode:  python ap5_tabela_descritiva_colesterol.py
4) O script imprime as tabelas no terminal e salva "ap5_tabela_descritiva_colesterol.xlsx"
   (abas "Numéricas" e "Categóricas") na mesma pasta do script.

Observações sobre os dados
--------------------------
- No CSV, as variáveis categóricas já vêm com o RÓTULO em texto (ex.: "Homem",
  "Sim") e as numéricas com o valor. Célula vazia = ausente (pergunta não aplicável
  ou sem resposta).
- Variáveis que são faixas/categorias codificadas com número (P02601, P04501,
  P04502) também funcionam se o seu CSV ainda tiver os códigos 1, 2, 3... em vez dos
  rótulos: o script converte usando o dicionário da PNS (campo "rotulos" abaixo).
- Q060 (a pergunta determinante) é constante ("Sim" para todos), então fica de fora
  das tabelas por padrão (veja IGNORAR). Tire-a de lá se quiser mostrá-la.
"""

import sys
from pathlib import Path

import numpy as np
import pandas as pd

# ---------------------------------------------------------------------------
# 0) CONFIGURAÇÃO
# ---------------------------------------------------------------------------

NOME_CSV = "pns_colesterol_25_59_v3.csv"
ARQUIVO_SAIDA = "ap5_tabela_descritiva_colesterol.xlsx"

INCLUIR_DIMENSAO = False     # True -> acrescenta a coluna "Dimensão" nas duas tabelas
MOSTRAR_PERCENTUAL = True    # True -> "Sim: 3229 (44,48%)"; False -> "Sim: 3229"
IGNORAR = {"Q060"}           # variáveis do CSV que NÃO entram nas tabelas
VALORES_AUSENTES_TEXTO = {"Ignorado"}   # rótulos que contam como ausente

try:
    PASTA = Path(__file__).resolve().parent
except NameError:            # rodando em célula interativa / notebook
    PASTA = Path.cwd()


# ---------------------------------------------------------------------------
# 1) ESPECIFICAÇÃO DAS VARIÁVEIS (67 variáveis + Q060)
#    classe: "numerica" ou "categorica"
#    tipo:   numéricas -> I (inteira) / R (real)
#            categóricas -> B (binária) / O (ordinal) / N (nominal)
#    rotulos: código PNS -> rótulo (só é usado se a coluna vier numérica no CSV;
#             a ordem dos rótulos também define a ordem natural das categorias)
#    ordem:   ordem natural das categorias de uma ordinal (quando não há "rotulos")
#    Categorias fora da lista "ordem" vão para o fim, da mais para a menos frequente.
# ---------------------------------------------------------------------------

NUM, CAT = "numerica", "categorica"


def V(codigo, nome, classe, tipo, dimensao, rotulos=None, ordem=None):
    return dict(codigo=codigo, nome=nome, classe=classe, tipo=tipo,
                dimensao=dimensao, rotulos=rotulos, ordem=ordem)


DEMO = "Demografia"
SAUDE = "Condições de saúde"
SOCIO = "Característica socioeconômica"
ALIM = "Hábitos alimentares"
VIDA = "Estilo de vida"

ROT_TV = {1: "Menos de uma hora", 2: "De uma hora a menos de duas horas",
          3: "De duas horas a menos de três horas", 4: "De três horas a menos de seis horas",
          5: "Seis horas ou mais", 6: "Não assiste televisão"}
ROT_TELA = {**ROT_TV, 6: "Não costuma usar computador, tablet ou celular no tempo livre"}
ROT_SAL = {1: "Muito alto", 2: "Alto", 3: "Adequado", 4: "Baixo", 5: "Muito baixo"}

VARSPEC = [
    # ---------------- Demografia (características do indivíduo + antropometria) ----------------
    V("C008", "Idade (anos)", NUM, "I", DEMO),
    V("C006", "Sexo", CAT, "B", DEMO),
    V("C009", "Cor ou raça", CAT, "N", DEMO),
    V("P00104", "Peso final (kg)", NUM, "R", DEMO),
    V("P00404", "Altura final (cm)", NUM, "I", DEMO),
    V("W00101", "Peso - 1ª pesagem (kg)", NUM, "R", DEMO),
    V("W00201", "Altura - 1ª medição (cm)", NUM, "R", DEMO),
    V("IMC", "Índice de Massa Corporal (calculado)", NUM, "R", DEMO),

    # ---------------- Condições de saúde ----------------
    V("Q00201", "Diagnóstico de hipertensão arterial", CAT, "B", SAUDE),
    V("Q03001", "Diagnóstico de diabetes", CAT, "B", SAUDE),
    V("Q068", "Diagnóstico de AVC", CAT, "B", SAUDE),
    V("Q074", "Diagnóstico de asma", CAT, "B", SAUDE),
    V("Q079", "Diagnóstico de artrite/reumatismo", CAT, "B", SAUDE),
    V("Q084", "Diagnóstico de problema crônico de coluna", CAT, "B", SAUDE),
    V("Q092", "Diagnóstico de depressão", CAT, "B", SAUDE),
    V("Q11604", "Diagnóstico de doença pulmonar crônica", CAT, "B", SAUDE),
    V("Q124", "Diagnóstico de insuficiência renal crônica", CAT, "B", SAUDE),
    V("Q120", "Diagnóstico de câncer", CAT, "B", SAUDE),
    V("N001", "Autoavaliação do estado de saúde", CAT, "O", SAUDE,
      ordem=["Muito boa", "Boa", "Regular", "Ruim", "Muito ruim"]),

    # ---------------- Característica socioeconômica ----------------
    V("V0001", "Unidade da Federação", CAT, "N", SOCIO),
    V("V0026", "Situação urbana/rural", CAT, "B", SOCIO),
    V("VDD004A", "Escolaridade (nível de instrução mais elevado)", CAT, "O", SOCIO,
      ordem=["Sem instrução", "Fundamental incompleto ou equivalente",
             "Fundamental completo ou equivalente", "Médio incompleto ou equivalente",
             "Médio completo ou equivalente", "Superior incompleto ou equivalente",
             "Superior completo"]),
    V("VDE001", "Condição na força de trabalho", CAT, "B", SOCIO),
    V("VDE002", "Condição de ocupação", CAT, "B", SOCIO),
    V("E001", "Trabalhou na semana de referência", CAT, "B", SOCIO),
    V("E01602", "Rendimento bruto mensal do trabalho (R$)", NUM, "I", SOCIO),
    V("VDF002", "Rendimento domiciliar (R$)", NUM, "I", SOCIO),
    V("VDF003", "Rendimento domiciliar per capita (R$)", NUM, "I", SOCIO),
    V("I00102", "Possui plano de saúde médico", CAT, "B", SOCIO),
    V("J001", "Estado de saúde (avaliação geral do morador)", CAT, "O", SOCIO,
      ordem=["Muito bom", "Bom", "Regular", "Ruim", "Muito ruim"]),
    V("J009", "Procura sempre o mesmo local/médico de saúde", CAT, "B", SOCIO),
    V("J01101", "Tempo desde a última consulta médica", CAT, "O", SOCIO,
      ordem=["Até 1 ano", "Mais de 1 ano a 2 anos", "Mais de 2 anos a 3 anos",
             "Mais de 3 anos", "Nunca foi ao médico"]),
    V("J012", "Nº de consultas médicas nos últimos 12 meses", NUM, "I", SOCIO),
    V("J014", "Procurou atendimento de saúde nas últimas 2 semanas", CAT, "B", SOCIO),

    # ---------------- Hábitos alimentares ----------------
    V("P006", "Dias/semana que come feijão (0-7)", NUM, "I", ALIM),
    V("P01101", "Dias/semana que come carne vermelha (0-7)", NUM, "I", ALIM),
    V("P013", "Dias/semana que come frango/galinha (0-7)", NUM, "I", ALIM),
    V("P015", "Dias/semana que come peixe (0-7)", NUM, "I", ALIM),
    V("P01001", "Vezes por dia que come verdura/legume", CAT, "O", ALIM,
      ordem=["Uma vez por dia (no almoço ou no jantar).",
             "Duas vezes por dia (no almoço e no jantar).",
             "Três vezes ou mais por dia."]),
    V("P01601", "Dias/semana que toma suco de fruta natural (0-7)", NUM, "I", ALIM),
    V("P018", "Dias/semana que come frutas (0-7)", NUM, "I", ALIM),
    V("P019", "Vezes por dia que come frutas", CAT, "O", ALIM,
      ordem=["Uma vez por dia", "Duas vezes por dia", "Três vezes ou mais por dia."]),
    V("P023", "Dias/semana que toma leite (0-7)", NUM, "I", ALIM),
    V("P02401", "Tipo de leite que costuma tomar", CAT, "N", ALIM),
    V("P02001", "Dias/semana que toma suco de caixinha/refresco em pó (0-7)", NUM, "I", ALIM),
    V("P02002", "Dias/semana que toma refrigerante (0-7)", NUM, "I", ALIM),
    V("P02501", "Dias/semana que come doces/ultraprocessados (0-7)", NUM, "I", ALIM),
    V("P02601", "Percepção do consumo de sal", CAT, "O", ALIM, rotulos=ROT_SAL),
    V("P02602", "Dias/semana que substitui o almoço por lanche rápido (0-7)", NUM, "I", ALIM),

    # ---------------- Estilo de vida ----------------
    V("P034", "Praticou exercício/esporte nos últimos 3 meses", CAT, "B", VIDA),
    V("P035", "Dias/semana que pratica exercício (0-7)", NUM, "I", VIDA),
    V("P036", "Exercício/esporte praticado com mais frequência", CAT, "N", VIDA),
    V("P03701", "Duração do exercício - horas", NUM, "I", VIDA),
    V("P03702", "Duração do exercício - minutos", NUM, "I", VIDA),
    V("P038", "Anda bastante a pé no trabalho", CAT, "B", VIDA),
    V("P04501", "Horas por dia assistindo TV (faixas)", CAT, "O", VIDA, rotulos=ROT_TV),
    V("P04502", "Horas por dia de tela no tempo livre (faixas)", CAT, "O", VIDA, rotulos=ROT_TELA),
    V("P050", "Fuma atualmente", CAT, "N", VIDA),
    V("P052", "Já fumou no passado", CAT, "N", VIDA),
    V("P05901", "Anos desde que parou de fumar", NUM, "I", VIDA),
    V("P027", "Frequência de consumo de bebida alcoólica", CAT, "O", VIDA,
      ordem=["Não bebo nunca", "Menos de uma vez por mês", "Uma vez ou mais por mês"]),
    V("P02801", "Dias/semana que consome álcool (0-7)", NUM, "I", VIDA),
    V("P029", "Doses de álcool no dia em que bebe", NUM, "I", VIDA),
    V("P03201", "Consumo de 5+ doses em uma ocasião (últimos 30 dias)", CAT, "B", VIDA),
    V("P03202", "Máximo de doses em uma única ocasião", NUM, "I", VIDA),
    V("Q060", "Diagnóstico de colesterol alto (determinante)", CAT, "B", VIDA),
    V("Q061", "Idade no diagnóstico de colesterol alto", NUM, "I", VIDA),
    V("Q06306", "Diagnóstico de doença do coração", CAT, "B", VIDA),
]


# ---------------------------------------------------------------------------
# 2) LEITURA DO CSV
# ---------------------------------------------------------------------------

def localizar_csv():
    """Procura o CSV: argumento da linha de comando > NOME_CSV > qualquer *colesterol*.csv."""
    if len(sys.argv) > 1:
        caminho = Path(sys.argv[1])
        if caminho.exists():
            return caminho
        raise SystemExit(f"Arquivo não encontrado: {caminho}")

    for pasta in (PASTA, Path.cwd()):
        for nome in (NOME_CSV, "pnscolesterol2559v3.csv"):
            if (pasta / nome).exists():
                return pasta / nome
        achados = sorted(pasta.glob("*colesterol*.csv"))
        if achados:
            return achados[0]

    csvs = [p.name for p in PASTA.glob("*.csv")]
    raise SystemExit(
        f"CSV '{NOME_CSV}' não encontrado em:\n  {PASTA}\n"
        f"CSVs que existem nessa pasta: {csvs or 'nenhum'}\n"
        "Coloque o CSV na mesma pasta do script ou ajuste NOME_CSV."
    )


def ler_csv(caminho):
    with open(caminho, "r", encoding="utf-8-sig", errors="ignore") as f:
        primeira = f.readline()
    sep = ";" if primeira.count(";") > primeira.count(",") else ","
    for enc in ("utf-8-sig", "latin1"):
        try:
            return pd.read_csv(caminho, sep=sep, encoding=enc, low_memory=False)
        except UnicodeDecodeError:
            continue
    raise SystemExit("Não consegui ler o CSV (encoding desconhecido).")


# ---------------------------------------------------------------------------
# 3) PREPARO DAS COLUNAS
# ---------------------------------------------------------------------------

def preparar_numerica(serie):
    return pd.to_numeric(serie, errors="coerce")


def preparar_categorica(serie, spec):
    """Devolve a coluna como texto (rótulos); ausentes viram NaN."""
    if pd.api.types.is_numeric_dtype(serie):
        # CSV com códigos numéricos: converte com o dicionário da PNS (campo "rotulos")
        rot = spec.get("rotulos") or {}
        s = serie.map(lambda x: np.nan if pd.isna(x) else rot.get(int(x), str(int(x))))
        s = s.astype(object)
    else:
        s = serie.astype(object).map(lambda x: np.nan if pd.isna(x) else str(x).strip())
        s = s.where(s != "", np.nan)
    return s.where(~s.isin(VALORES_AUSENTES_TEXTO), np.nan)


def _valor(x):
    """Inteiro quando o número é redondo; senão, 2 casas decimais."""
    if pd.isna(x):
        return np.nan
    x = float(x)
    return int(x) if x.is_integer() else round(x, 2)


def _pct(p):
    return f"{p:.2f}".replace(".", ",") + "%"


# ---------------------------------------------------------------------------
# 4) TABELA DE VARIÁVEIS NUMÉRICAS
# ---------------------------------------------------------------------------

def tabela_numericas(df, specs):
    linhas = []
    for spec in specs:
        if spec["classe"] != NUM:
            continue
        bruta = df[spec["codigo"]]
        s = preparar_numerica(bruta)
        perdidos = int(bruta.notna().sum() - s.notna().sum())
        if perdidos:
            print(f"[AVISO] {spec['codigo']}: {perdidos} valor(es) não numérico(s) "
                  f"foram tratados como ausentes.")
        tem = s.notna().any()
        linha = {
            "Variável": f'{spec["codigo"]} - {spec["nome"]}',
            "Tipo": spec["tipo"],
            "Nº de Ausentes": int(s.isna().sum()),
            "Mínimo": _valor(s.min()) if tem else np.nan,
            "Máximo": _valor(s.max()) if tem else np.nan,
            "Média": round(s.mean(), 2) if tem else np.nan,
            "Mediana": _valor(s.median()) if tem else np.nan,
            "Desvio Padrão": round(s.std(), 2) if tem else np.nan,
        }
        if INCLUIR_DIMENSAO:
            linha = {"Dimensão": spec["dimensao"], **linha}
        linhas.append(linha)
    return pd.DataFrame(linhas)


# ---------------------------------------------------------------------------
# 5) TABELA DE VARIÁVEIS CATEGÓRICAS
# ---------------------------------------------------------------------------

def tabela_categoricas(df, specs):
    linhas = []
    for spec in specs:
        if spec["classe"] != CAT:
            continue
        s = preparar_categorica(df[spec["codigo"]], spec)
        cont = s.value_counts(dropna=True)          # já vem do mais para o menos frequente
        validos = int(cont.sum())

        ordem = spec.get("ordem") or (list(spec["rotulos"].values()) if spec.get("rotulos") else [])
        valores = [v for v in ordem if v in cont.index] + [v for v in cont.index if v not in ordem]

        if spec["tipo"] == "B" and len(valores) != 2:
            print(f"[AVISO] {spec['codigo']} foi declarada Binária (B), mas tem "
                  f"{len(valores)} valores distintos. Revise o tipo no VARSPEC.")

        def freq(v):
            n = int(cont[v])
            return f"{v}: {n} ({_pct(100 * n / validos)})" if MOSTRAR_PERCENTUAL else f"{v}: {n}"

        linha = {
            "Variável": f'{spec["codigo"]} - {spec["nome"]}',
            "Tipo": spec["tipo"],
            "Nº de Ausentes": int(s.isna().sum()),
            "Valores": "; ".join(valores),
            "Frequência por valor": "; ".join(freq(v) for v in valores),
            "Moda": cont.index[0] if validos else "—",
        }
        if INCLUIR_DIMENSAO:
            linha = {"Dimensão": spec["dimensao"], **linha}
        linhas.append(linha)
    return pd.DataFrame(linhas)


# ---------------------------------------------------------------------------
# 6) SAÍDA (EXCEL)
# ---------------------------------------------------------------------------

def salvar_excel(tab_num, tab_cat, caminho):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        tab_num.to_excel(writer, sheet_name="Numéricas", index=False)
        tab_cat.to_excel(writer, sheet_name="Categóricas", index=False)
        for ws in writer.book.worksheets:
            for cel in ws[1]:
                cel.font = Font(bold=True, color="FFFFFF")
                cel.fill = PatternFill("solid", start_color="1F4E78")
                cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col in ws.columns:
                letra = get_column_letter(col[0].column)
                maior = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[letra].width = min(max(maior + 2, 10), 70)
                for c in col[1:]:
                    c.alignment = Alignment(vertical="top", wrap_text=True)
            ws.freeze_panes = "A2"


def salvar(tab_num, tab_cat):
    destino = PASTA / ARQUIVO_SAIDA
    try:
        try:
            salvar_excel(tab_num, tab_cat, destino)
        except PermissionError:   # arquivo aberto no Excel
            destino = destino.with_name(destino.stem + "_novo.xlsx")
            salvar_excel(tab_num, tab_cat, destino)
        print(f"\nArquivo salvo em: {destino}")
    except ImportError:
        tab_num.to_csv(PASTA / "ap5_tabela_numericas.csv", index=False, encoding="utf-8-sig")
        tab_cat.to_csv(PASTA / "ap5_tabela_categoricas.csv", index=False, encoding="utf-8-sig")
        print("\n[AVISO] 'openpyxl' não está instalado; salvei em CSV na pasta do script.")
        print("Para gerar o .xlsx, rode no terminal:  pip install openpyxl")


# ---------------------------------------------------------------------------
# 7) EXECUÇÃO PRINCIPAL
# ---------------------------------------------------------------------------

def main():
    caminho = localizar_csv()
    print(f"Lendo: {caminho}")
    df = ler_csv(caminho)
    print(f"Registros: {len(df)} | Colunas no CSV: {len(df.columns)}")

    specs = [s for s in VARSPEC if s["codigo"] not in IGNORAR]

    faltando = [s["codigo"] for s in specs if s["codigo"] not in df.columns]
    if faltando:
        print(f"[AVISO] Variáveis do VARSPEC que NÃO estão no CSV (serão puladas): {faltando}")
        specs = [s for s in specs if s["codigo"] in df.columns]

    extras = [c for c in df.columns if c not in {s["codigo"] for s in VARSPEC}]
    if extras:
        print(f"[AVISO] Colunas do CSV que não estão no VARSPEC (não entram nas tabelas): {extras}")

    if len(df) == 0:
        raise SystemExit("O CSV está vazio.")

    tab_num = tabela_numericas(df, specs)
    tab_cat = tabela_categoricas(df, specs)

    pd.set_option("display.width", 250)
    print(f"\n===== VARIÁVEIS NUMÉRICAS ({len(tab_num)}) =====")
    print(tab_num.to_string(index=False, max_colwidth=60))
    print(f"\n===== VARIÁVEIS CATEGÓRICAS ({len(tab_cat)}) =====")
    print(tab_cat.to_string(index=False, max_colwidth=60))
    print(f"\nTotal de variáveis nas tabelas: {len(tab_num) + len(tab_cat)}")

    salvar(tab_num, tab_cat)


if __name__ == "__main__":
    main()